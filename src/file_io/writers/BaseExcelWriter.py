import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

class BaseExcelWriter:
    """
    Infrastructure component to write tabular data to a real Excel (.xlsx) file,
    allowing automatic column width configuration and native table styling.
    """
    
    @staticmethod
    def write(path: str, headers: list[str], rows: list[list[str]]) -> None:
        if not path:
            return
            
        # Create a new workbook and select active sheet
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Schedule"
        
        # Append headers and all rows
        ws.append(headers)
        for row in rows:
            ws.append(row)
        
        # Calculate the exact range of our data (e.g., A1:E25)
        max_row = len(rows) + 1
        max_col = len(headers)
        max_col_letter = get_column_letter(max_col)
        table_ref = f"A1:{max_col_letter}{max_row}"

        # Create the Native Excel Table
        tab = Table(displayName="ExamSchedule", ref=table_ref)

        # Apply the requested style: "TableStyleLight20" (Light blue with zebra striping)
        style = TableStyleInfo(
            name="TableStyleLight20", 
            showFirstColumn=False,
            showLastColumn=False, 
            # Zebra striping
            showRowStripes=True,
            showColumnStripes=False
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)

        # Auto-fit column widths dynamically based on the longest text in each column
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            # Add a padding of 4 characters so it looks spacious and neat
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12) 
            
        # Save the workbook to the targeted path
        wb.save(path)